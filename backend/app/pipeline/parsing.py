"""Extração de texto de contratos PDF/DOCX (primeiro estágio do pipeline)."""

import logging
from pathlib import Path

import docx
import pdfplumber

logger = logging.getLogger(__name__)

MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

MIN_TEXT_CHARS = 50
MAX_TEXT_CHARS = 150_000


class ParsingError(Exception):
    """Erro na extração de texto de um documento."""


def parse_document(path: Path, mime_type: str, max_chars: int = MAX_TEXT_CHARS) -> str:
    """Extrai o texto de um documento PDF ou DOCX.

    Raises:
        ParsingError: mime_type não suportado ou documento sem texto extraível.
    """
    if mime_type == MIME_PDF:
        text = _parse_pdf(path)
        if len(text.strip()) < MIN_TEXT_CHARS:
            raise ParsingError("PDF sem texto extraível (OCR fora de escopo)")
    elif mime_type == MIME_DOCX:
        text = _parse_docx(path)
        if len(text.strip()) < MIN_TEXT_CHARS:
            raise ParsingError("DOCX sem texto extraível (OCR fora de escopo)")
    else:
        raise ParsingError(f"Tipo de arquivo não suportado: {mime_type}")

    if len(text) > max_chars:
        logger.warning(
            "Texto truncado de %d para %d caracteres: %s", len(text), max_chars, path.name
        )
        text = text[:max_chars]
    return text


def _parse_pdf(path: Path) -> str:
    with pdfplumber.open(path) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages)


def _parse_docx(path: Path) -> str:
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
